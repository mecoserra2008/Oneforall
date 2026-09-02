"""The Excel deliverable.

Twelve sheets, in the order a reader needs them:

    Dashboard      the answer, on one screen
    Bridge         where the day's PnL came from, with the tie-outs
    Snail          the snail trail: today walked leg by leg, and the history
    Asset Classes  bonds, futures and swaps side by side
    Bonds          every attribution column, one row per position
    Futures        one row per contract
    Swaps          one row per deal
    Frameworks     how the book is being measured, and what each chain is
    Breakdowns     by portfolio and by currency
    Curves         the curves everything was calibrated on
    Data Quality   what the report could not see
    Definitions    every column, its unit, its meaning and its formula

Two decisions worth stating.

VALUES, NOT FORMULAS.  The legacy workbook wrote Excel formulas so the sheet
recalculated itself.  That made sense when Excel WAS the engine.  It is not any
more: the engine is Python, the maths is tested, and a sheet that recomputes the
model in a second language is a second model that can disagree with the first.
What is preserved instead is auditability - every figure traces to a column on
the Bonds sheet, every column is defined on the Definitions sheet, and the CSV
and JSON outputs carry the same numbers for anything downstream.

CHARTS ARE NEVER FATAL.  A chart that fails to build must not cost the reader
the report.  Each is wrapped, and a failure leaves a note in its place.
"""

from __future__ import annotations

import datetime as _dt
from pathlib import Path
from typing import Any, Iterable, Mapping, Sequence

import numpy as np
from openpyxl import Workbook
from openpyxl.chart import BarChart, Reference, ScatterChart, Series
from openpyxl.chart.data_source import NumDataSource, NumRef
from openpyxl.chart.marker import Marker
from openpyxl.chart.shapes import GraphicalProperties
from openpyxl.chart.trendline import Trendline
from openpyxl.drawing.line import LineProperties
from openpyxl.formatting.rule import CellIsRule, ColorScaleRule, DataBarRule, FormulaRule
from openpyxl.utils import get_column_letter
from openpyxl.worksheet.worksheet import Worksheet

from ..aggregate import PortfolioSummary
from ..columns import Column, ColumnFrame, Kind
from ..engine import AttributionResult
from ..snail import SnailTrail
from .theme import (
    CENTRE,
    COLUMN_WIDTH,
    LEFT,
    PALETTE,
    RIGHT,
    WRAP_LEFT,
    body_font,
    bottom_rule,
    fill,
    header_font,
    number_format_for,
    thin_border,
    title_font,
)

__all__ = ["write_workbook"]

_EUR = number_format_for(Kind.EUR)
_EUR_C = number_format_for(Kind.EUR, cents=True)
_PCT = number_format_for(Kind.PCT)


# --------------------------------------------------------------------------- #
# small drawing helpers
# --------------------------------------------------------------------------- #


def _banner(ws: Worksheet, title: str, subtitle: str, width: int = 12) -> int:
    """The dark title bar every sheet opens with.  Returns the next free row."""
    ws.merge_cells(start_row=1, start_column=1, end_row=2, end_column=width)
    cell = ws.cell(row=1, column=1, value=title)
    cell.font = title_font(18)
    cell.alignment = LEFT
    for col in range(1, width + 1):
        for row in (1, 2):
            ws.cell(row=row, column=col).fill = fill(PALETTE.ink)

    ws.merge_cells(start_row=3, start_column=1, end_row=3, end_column=width)
    sub = ws.cell(row=3, column=1, value=subtitle)
    sub.font = body_font(9, colour=PALETTE.surface)
    sub.alignment = LEFT
    for col in range(1, width + 1):
        ws.cell(row=3, column=col).fill = fill(PALETTE.ink_soft)

    ws.row_dimensions[1].height = 24
    ws.row_dimensions[2].height = 12
    ws.row_dimensions[3].height = 16
    return 5


def _section(ws: Worksheet, row: int, title: str, note: str = "", width: int = 12) -> int:
    """A section heading with an optional one-line explanation beneath it."""
    cell = ws.cell(row=row, column=1, value=title)
    cell.font = body_font(12, bold=True, colour=PALETTE.ink)
    for col in range(1, width + 1):
        ws.cell(row=row, column=col).border = bottom_rule(PALETTE.ink, "medium")
    row += 1
    if note:
        ws.merge_cells(start_row=row, start_column=1, end_row=row, end_column=width)
        n = ws.cell(row=row, column=1, value=note)
        n.font = body_font(9, italic=True, colour=PALETTE.muted)
        n.alignment = WRAP_LEFT
        ws.row_dimensions[row].height = 26
        row += 1
    return row + 1


def _tile(
    ws: Worksheet,
    row: int,
    col: int,
    label: str,
    value: Any,
    fmt: str = _EUR,
    *,
    span: int = 2,
    tone: str | None = None,
    caption: str = "",
) -> None:
    """A KPI tile: caption above, big number below, on a soft panel."""
    ws.merge_cells(start_row=row, start_column=col, end_row=row, end_column=col + span - 1)
    head = ws.cell(row=row, column=col, value=label)
    head.font = body_font(9, bold=True, colour=PALETTE.muted)
    head.alignment = CENTRE

    ws.merge_cells(start_row=row + 1, start_column=col, end_row=row + 1, end_column=col + span - 1)
    body = ws.cell(row=row + 1, column=col, value=_clean(value))
    body.font = body_font(16, bold=True, colour=tone or PALETTE.ink)
    body.alignment = CENTRE
    body.number_format = fmt

    ws.merge_cells(start_row=row + 2, start_column=col, end_row=row + 2, end_column=col + span - 1)
    foot = ws.cell(row=row + 2, column=col, value=caption)
    foot.font = body_font(8, italic=True, colour=PALETTE.muted)
    foot.alignment = CENTRE

    for r in range(row, row + 3):
        for c in range(col, col + span):
            ws.cell(row=r, column=c).fill = fill(PALETTE.panel)
            ws.cell(row=r, column=c).border = thin_border()
    ws.row_dimensions[row].height = 15
    ws.row_dimensions[row + 1].height = 24
    ws.row_dimensions[row + 2].height = 13


def _table_header(ws: Worksheet, row: int, headers: Sequence[str], start_col: int = 1) -> None:
    for i, text in enumerate(headers):
        cell = ws.cell(row=row, column=start_col + i, value=text)
        cell.font = header_font(10)
        cell.fill = fill(PALETTE.ink)
        cell.alignment = CENTRE
        cell.border = thin_border(PALETTE.ink)
    ws.row_dimensions[row].height = 28


def _clean(value: Any) -> Any:
    """openpyxl cannot store NaN or a numpy scalar; blank is the honest form."""
    if value is None:
        return None
    if isinstance(value, (np.floating, float)):
        f = float(value)
        return None if (f != f or f in (float("inf"), float("-inf"))) else f
    if isinstance(value, (np.integer,)):
        return int(value)
    if isinstance(value, (np.bool_,)):
        return bool(value)
    if isinstance(value, np.datetime64):
        return None if np.isnat(value) else value.astype("datetime64[D]").astype(_dt.date)
    if isinstance(value, np.str_):
        return str(value)
    return value


def _tone(value: Any) -> str:
    """Green for a gain, red for a loss - applied only where sign is the point."""
    v = _clean(value)
    if not isinstance(v, (int, float)):
        return PALETTE.ink
    return PALETTE.positive if v >= 0 else PALETTE.negative


def _widths(ws: Worksheet, widths: Mapping[int, int]) -> None:
    for col, width in widths.items():
        ws.column_dimensions[get_column_letter(col)].width = width


def _no_line() -> GraphicalProperties:
    return GraphicalProperties(ln=LineProperties(noFill=True))


def _solid(colour: str) -> GraphicalProperties:
    return GraphicalProperties(solidFill=colour)


def _safe_chart(builder, ws: Worksheet, anchor: str, label: str) -> None:
    """Build a chart, or leave a note where it would have been.

    A chart is decoration; the tables are the report.  Losing the whole file
    because one series would not build is not a trade worth making.
    """
    try:
        chart = builder()
        if chart is not None:
            ws.add_chart(chart, anchor)
    except Exception as exc:  # pragma: no cover - defensive
        cell = ws[anchor]
        cell.value = f"[chart unavailable: {label} - {type(exc).__name__}]"
        cell.font = body_font(9, italic=True, colour=PALETTE.muted)


# --------------------------------------------------------------------------- #
# generic frame sheet
# --------------------------------------------------------------------------- #


def _write_frame_sheet(
    ws: Worksheet,
    frame: ColumnFrame,
    title: str,
    subtitle: str,
    *,
    freeze_cols: int = 2,
    highlight: Mapping[str, str] | None = None,
) -> None:
    """One sheet per position table: banner, grouped header, formatted rows."""
    n_cols = len(frame.schema)
    row = _banner(ws, title, subtitle, width=min(n_cols, 14))

    # Group strip above the header, so 90 columns are navigable.
    group_row = row
    col = 1
    while col <= n_cols:
        group = frame.schema[col - 1].group or " "
        span = 1
        while col + span <= n_cols and (frame.schema[col + span - 1].group or " ") == group:
            span += 1
        ws.merge_cells(start_row=group_row, start_column=col, end_row=group_row, end_column=col + span - 1)
        cell = ws.cell(row=group_row, column=col, value=group.upper())
        cell.font = body_font(9, bold=True, colour=PALETTE.surface)
        cell.alignment = CENTRE
        for c in range(col, col + span):
            ws.cell(row=group_row, column=c).fill = fill(PALETTE.accent)
            ws.cell(row=group_row, column=c).border = thin_border(PALETTE.accent)
        col += span

    header_row = group_row + 1
    _table_header(ws, header_row, [c.label for c in frame.schema])

    # Hovering the header gives the economic meaning, so the definition travels
    # with the number instead of living on another sheet.
    for i, column in enumerate(frame.schema, start=1):
        cell = ws.cell(row=header_row, column=i)
        note = column.meaning
        if column.formula:
            note = f"{note}\n\n{column.formula}"
        try:
            from openpyxl.comments import Comment

            cell.comment = Comment(note, "pnlx", width=320, height=140)
        except Exception:  # pragma: no cover - comments are a nicety
            pass

    first_data = header_row + 1
    for r, values in enumerate(frame.rows(), start=first_data):
        banded = (r - first_data) % 2 == 1
        for i, (column, value) in enumerate(zip(frame.schema, values), start=1):
            cell = ws.cell(row=r, column=i, value=_clean(value))
            cell.number_format = number_format_for(column.kind)
            cell.font = body_font(9)
            cell.alignment = LEFT if column.kind in (Kind.TEXT, Kind.FLAG) else RIGHT
            if banded:
                cell.fill = fill(PALETTE.panel)

    last_row = first_data + len(frame) - 1
    if len(frame) == 0:
        ws.cell(row=first_data, column=1, value="no positions").font = body_font(
            9, italic=True, colour=PALETTE.muted
        )
        return

    ws.freeze_panes = ws.cell(row=first_data, column=freeze_cols + 1)
    ws.auto_filter.ref = f"A{header_row}:{get_column_letter(n_cols)}{last_row}"

    for i, column in enumerate(frame.schema, start=1):
        ws.column_dimensions[get_column_letter(i)].width = COLUMN_WIDTH.get(column.kind, 13)

    # Conditional formatting, only where it says something.
    names = {c.name: i + 1 for i, c in enumerate(frame.schema)}
    for column_name, style in (highlight or {}).items():
        idx = names.get(column_name)
        if idx is None:
            continue
        letter = get_column_letter(idx)
        rng = f"{letter}{first_data}:{letter}{last_row}"
        if style == "diverging":
            ws.conditional_formatting.add(
                rng,
                ColorScaleRule(
                    start_type="percentile", start_value=5, start_color="F5C6C2",
                    mid_type="num", mid_value=0, mid_color="FFFFFF",
                    end_type="percentile", end_value=95, end_color="BFE3CE",
                ),
            )
        elif style == "bar":
            ws.conditional_formatting.add(
                rng, DataBarRule(start_type="min", end_type="max", color=PALETTE.accent)
            )
        elif style == "status":
            ws.conditional_formatting.add(
                rng,
                CellIsRule(
                    operator="notEqual", formula=['"OK"'],
                    font=body_font(9, bold=True, colour=PALETTE.negative),
                    fill=fill("FBE9E7"),
                ),
            )


# --------------------------------------------------------------------------- #
# dashboard
# --------------------------------------------------------------------------- #


def _dashboard(ws: Worksheet, result: AttributionResult, summary: PortfolioSummary, trail: SnailTrail) -> None:
    cfg = result.config
    h = summary.headline
    q = summary.quality

    row = _banner(
        ws,
        "PnL Explainer - Desk Overview",
        f"{cfg.run.book}   |   {cfg.run.prior:%d %b %Y}  to  {cfg.run.as_of:%d %b %Y}   "
        f"({cfg.run.days} day{'s' if cfg.run.days != 1 else ''})   |   base {cfg.run.base_ccy}   "
        f"|   built {_dt.datetime.now():%d %b %Y %H:%M}",
        width=12,
    )

    # ---- headline tiles ----------------------------------------------------- #
    row = _section(
        ws, row, "The day in one line",
        "Practical is what the book made from marks and cash. Theoretical is what the risk "
        "model says it should have made. Explained adds the hedge basis, so the residual "
        "below is bond-leg model error and nothing else.",
    )

    _tile(ws, row, 1, "PRACTICAL PnL", h["practical_pnl"], _EUR,
          tone=_tone(h["practical_pnl"]), caption="marks + coupon + actual hedges")
    _tile(ws, row, 3, "THEORETICAL PnL", h["theoretical_pnl"], _EUR,
          tone=_tone(h["theoretical_pnl"]), caption="what the model says")
    _tile(ws, row, 5, "TOTAL EXPLAINED", h["total_explained"], _EUR,
          tone=_tone(h["total_explained"]), caption="theoretical + hedge basis")
    _tile(ws, row, 7, "RESIDUAL", h["residual_pnl"], _EUR_C,
          tone=_tone(-abs(h["residual_pnl"])) if abs(h["residual_pnl"]) > 0 else PALETTE.ink,
          caption="practical less explained")
    _tile(ws, row, 9, "RESIDUAL % OF PRACTICAL", h["residual_pct"], _PCT,
          caption="how much is unexplained")
    _tile(ws, row, 11, "BRIDGE TIE-OUT", q.get("bridge_tie_out", 0.0), '0.000000',
          caption="must be zero")
    row += 4

    _tile(ws, row, 1, "DURATION", h["duration_total"], _EUR, tone=_tone(h["duration_total"]),
          caption="rates and spread, first order")
    _tile(ws, row, 3, "CONVEXITY", h["convexity"], _EUR, tone=_tone(h["convexity"]),
          caption="second order in the same move")
    _tile(ws, row, 5, "CARRY", h["carry_total"], _EUR, tone=_tone(h["carry_total"]),
          caption="coupon + pull to par")
    _tile(ws, row, 7, "FX", h["pnl_fx"], _EUR, tone=_tone(h["pnl_fx"]),
          caption="translation + cross term")
    _tile(ws, row, 9, "HEDGE (MODEL)", h["model_hedge_pnl"], _EUR, tone=_tone(h["model_hedge_pnl"]),
          caption="futures + swaps vs their curves")
    _tile(ws, row, 11, "HEDGE BASIS", h["hedge_basis_pnl"], _EUR, tone=_tone(h["hedge_basis_pnl"]),
          caption="actual less model hedge")
    row += 4

    _tile(ws, row, 1, "BOND DV01", h["bond_dv01"], _EUR, caption="risk at the close")
    _tile(ws, row, 3, "HEDGE DV01", h["hedge_dv01"], _EUR, caption="futures + plain swaps")
    _tile(ws, row, 5, "RESIDUAL DV01", h["residual_dv01"], _EUR,
          tone=_tone(-abs(h["residual_dv01"])), caption="net risk; zero is flat")
    _tile(ws, row, 7, "HEDGE EFFICIENCY", h["portfolio_hedge_efficiency"], _PCT,
          caption="netted across the book")
    _tile(ws, row, 9, "MEDIAN PER-BOND EFF.", h["median_hedge_efficiency"], _PCT,
          caption="the gap to the left is the story")
    _tile(ws, row, 11, "UNLINKED HEDGE PnL", h["unlinked_hedge_pnl"], _EUR,
          tone=PALETTE.warn if abs(h["unlinked_hedge_pnl"]) > 0 else PALETTE.ink,
          caption="real PnL on no bond row")
    row += 5

    # ---- bridge + waterfall -------------------------------------------------- #
    row = _section(
        ws, row, "Where the PnL came from",
        "Cumulative walk from zero to the explained total. Each bar starts where the last "
        "one finished, so the height of a bar is that leg's contribution and the end of the "
        "run is the total.",
    )
    bridge_top = row
    _table_header(ws, row, ["Component", "PnL", "% of practical"])
    row += 1
    denom = abs(h["practical_pnl"]) or float("nan")
    for line in summary.bridge.lines:
        if line.kind == "check":
            continue
        label = ("    " + line.label) if line.is_memo else line.label
        c1 = ws.cell(row=row, column=1, value=label)
        c2 = ws.cell(row=row, column=2, value=_clean(line.amount))
        c3 = ws.cell(row=row, column=3, value=_clean(line.amount / denom if denom == denom else None))
        c2.number_format = _EUR
        c3.number_format = _PCT
        if line.is_memo:
            for c in (c1, c2, c3):
                c.font = body_font(9, italic=True, colour=PALETTE.memo)
        elif line.kind in ("subtotal", "total"):
            for c in (c1, c2, c3):
                c.font = body_font(10, bold=True, colour=PALETTE.ink)
                c.fill = fill(PALETTE.accent_soft)
        else:
            c1.font = body_font(10)
            c2.font = body_font(10, colour=_tone(line.amount))
            c3.font = body_font(10)
        for c in (c1, c2, c3):
            c.border = thin_border()
        row += 1
    bridge_bottom = row - 1

    _widths(ws, {1: 46, 2: 16, 3: 14})
    for col in range(4, 13):
        ws.column_dimensions[get_column_letter(col)].width = 13

    _safe_chart(
        lambda: _waterfall_chart(ws, summary, anchor_row=bridge_top),
        ws, f"E{bridge_top}", "bridge waterfall",
    )

    row += 2

    # ---- asset classes ------------------------------------------------------- #
    row = _section(
        ws, row, "By instrument type",
        "Each class on its own terms. Bonds are the cash leg alone - hedges appear once, "
        "in their own class - so the residual on the bond line is the same figure as the "
        "bridge's unexplained residual.",
    )
    ac_top = row
    _table_header(
        ws, row,
        ["Instrument type", "Positions", "DV01", "Theoretical", "Practical", "Difference", "Coverage"],
    )
    row += 1
    for ac in summary.asset_classes:
        values = [ac.name, ac.count, ac.dv01, ac.theoretical, ac.practical, ac.residual, ac.coverage]
        formats = ["@", "0", _EUR, _EUR, _EUR, _EUR_C, _PCT]
        for i, (value, fmt) in enumerate(zip(values, formats), start=1):
            cell = ws.cell(row=row, column=i, value=_clean(value))
            cell.number_format = fmt
            cell.font = body_font(10, colour=_tone(value) if i in (4, 5, 6) else PALETTE.text)
            cell.border = thin_border()
        row += 1
    ac_bottom = row - 1

    _safe_chart(
        lambda: _asset_class_chart(ws, ac_top, ac_bottom),
        ws, f"I{ac_top}", "asset class comparison",
    )
    row += 2

    # ---- snail --------------------------------------------------------------- #
    row = _section(
        ws, row, "Snail trail - today, walked leg by leg",
        "Each step adds one leg and plots what is still outstanding. A leg that barely moves "
        "the line is not explaining anything; a long run means that leg is doing the work. "
        "The walk closes on the practical total by construction.",
    )
    snail_top = row
    _table_header(ws, row, ["#", "Leg", "Amount", "Cumulative explained", "Still outstanding"])
    row += 1
    for step in trail.steps:
        values = [step.order, step.label, step.amount, step.cumulative_explained, step.outstanding]
        formats = ["0", "@", _EUR, _EUR, _EUR]
        for i, (value, fmt) in enumerate(zip(values, formats), start=1):
            cell = ws.cell(row=row, column=i, value=_clean(value))
            cell.number_format = fmt
            cell.font = body_font(10, bold=(step.label == "Residual"))
            cell.border = thin_border()
        row += 1
    snail_bottom = row - 1

    _safe_chart(
        lambda: _snail_step_chart(ws, snail_top, snail_bottom),
        ws, f"G{snail_top}", "intraday snail",
    )
    row += 2

    # ---- top movers ---------------------------------------------------------- #
    row = _section(
        ws, row, "Largest unexplained residuals",
        "Ranked on absolute residual, not percentage: a 40% break on a 60 EUR position is "
        "arithmetic, a 2% break on a 2m position is a problem.",
    )
    top = _top_by_absolute(result.bonds, "residual_pnl", cfg.output.top_n)
    _table_header(ws, row, ["ISIN", "Name", "Framework", "Practical", "Explained", "Residual", "Residual %", "Status"])
    row += 1
    for record in top:
        values = [
            record["isin"], record["name"], record["spread_framework"],
            record["practical_pnl"], record["total_explained"],
            record["residual_pnl"], record["residual_pct"], record["attribution_status"],
        ]
        formats = ["@", "@", "@", _EUR, _EUR, _EUR_C, _PCT, "@"]
        for i, (value, fmt) in enumerate(zip(values, formats), start=1):
            cell = ws.cell(row=row, column=i, value=_clean(value))
            cell.number_format = fmt
            cell.font = body_font(9)
            cell.border = thin_border()
            if i == 8 and str(value) != "OK":
                cell.font = body_font(9, bold=True, colour=PALETTE.negative)
        row += 1

    ws.sheet_view.showGridLines = False


def _top_by_absolute(frame: ColumnFrame, column: str, limit: int) -> list[dict[str, Any]]:
    if len(frame) == 0:
        return []
    values = np.abs(np.nan_to_num(np.asarray(frame.get(column), dtype=np.float64)))
    order = np.argsort(values)[::-1][:limit]
    records = frame.to_records()
    return [records[i] for i in order.tolist() if values[i] > 0]


# --------------------------------------------------------------------------- #
# charts
# --------------------------------------------------------------------------- #


def _waterfall_chart(ws: Worksheet, summary: PortfolioSummary, anchor_row: int):
    """A waterfall, built as a stacked bar with an invisible base series.

    Excel has no waterfall type openpyxl can emit, so the standard construction
    is used: an invisible spacer that lifts each bar to where the previous one
    finished, then the visible bar on top of it.  Rises and falls are separate
    series so they can be coloured differently - which is the whole point of a
    waterfall.
    """
    steps = summary.bridge.waterfall()
    if not steps:
        return None

    # Staging block, off to the right and hidden: charts need cells to read.
    first = anchor_row
    col_label, col_base, col_rise, col_fall = 20, 21, 22, 23
    ws.cell(row=first - 1, column=col_label, value="waterfall")
    for i, (label, start, end, delta) in enumerate(steps):
        r = first + i
        ws.cell(row=r, column=col_label, value=label)
        ws.cell(row=r, column=col_base, value=min(start, end))
        ws.cell(row=r, column=col_rise, value=abs(delta) if delta >= 0 else 0)
        ws.cell(row=r, column=col_fall, value=abs(delta) if delta < 0 else 0)
    last = first + len(steps) - 1

    for c in (col_label, col_base, col_rise, col_fall):
        ws.column_dimensions[get_column_letter(c)].hidden = True

    chart = BarChart()
    chart.type = "col"
    chart.grouping = "stacked"
    chart.overlap = 100
    chart.title = "PnL bridge"
    chart.height = 9.5
    chart.width = 20
    chart.gapWidth = 40

    base = Series(Reference(ws, min_col=col_base, min_row=first, max_row=last), title="base")
    base.graphicalProperties = GraphicalProperties(noFill=True, ln=LineProperties(noFill=True))
    rise = Series(Reference(ws, min_col=col_rise, min_row=first, max_row=last), title="gain")
    rise.graphicalProperties = _solid(PALETTE.positive)
    fall = Series(Reference(ws, min_col=col_fall, min_row=first, max_row=last), title="loss")
    fall.graphicalProperties = _solid(PALETTE.negative)

    for series in (base, rise, fall):
        chart.series.append(series)

    chart.set_categories(Reference(ws, min_col=col_label, min_row=first, max_row=last))
    chart.y_axis.title = "EUR"
    chart.y_axis.numFmt = "#,##0"
    chart.legend = None
    return chart


def _asset_class_chart(ws: Worksheet, top: int, bottom: int):
    """Theoretical against practical, one pair of bars per instrument type."""
    chart = BarChart()
    chart.type = "col"
    chart.grouping = "clustered"
    chart.title = "Theoretical vs practical by instrument type"
    chart.height = 8
    chart.width = 14
    data = Reference(ws, min_col=4, max_col=5, min_row=top, max_row=bottom)
    chart.add_data(data, titles_from_data=True)
    chart.set_categories(Reference(ws, min_col=1, min_row=top + 1, max_row=bottom))
    chart.series[0].graphicalProperties = _solid(PALETTE.accent)
    chart.series[1].graphicalProperties = _solid(PALETTE.ink)
    chart.y_axis.numFmt = "#,##0"
    return chart


def _snail_step_chart(ws: Worksheet, top: int, bottom: int):
    """The intraday walk: cumulative explained against what is still outstanding."""
    chart = BarChart()
    chart.type = "col"
    chart.grouping = "clustered"
    chart.title = "Snail trail - cumulative explanation"
    chart.height = 8
    chart.width = 14
    data = Reference(ws, min_col=4, max_col=5, min_row=top, max_row=bottom)
    chart.add_data(data, titles_from_data=True)
    chart.set_categories(Reference(ws, min_col=2, min_row=top + 1, max_row=bottom))
    chart.series[0].graphicalProperties = _solid(PALETTE.accent)
    chart.series[1].graphicalProperties = _solid(PALETTE.warn)
    chart.y_axis.numFmt = "#,##0"
    return chart


def _history_snail_chart(ws: Worksheet, top: int, bottom: int, x_col: int, y_col: int, title: str):
    """The historical trail as a connected scatter - the snail proper.

    A perfect model traces the 45-degree line; the distance from it is the
    accumulated unexplained PnL and the SHAPE says what kind of error it is.
    """
    chart = ScatterChart()
    chart.title = title
    chart.style = 13
    chart.height = 10
    chart.width = 18
    chart.x_axis.title = "Cumulative practical PnL"
    chart.y_axis.title = "Cumulative explained PnL"
    chart.x_axis.numFmt = "#,##0"
    chart.y_axis.numFmt = "#,##0"
    # Both axes must stay visible or the plot reads as a blank box in Excel.
    chart.x_axis.delete = False
    chart.y_axis.delete = False

    xvalues = Reference(ws, min_col=x_col, min_row=top, max_row=bottom)
    yvalues = Reference(ws, min_col=y_col, min_row=top, max_row=bottom)
    series = Series(yvalues, xvalues, title="trail")
    series.marker = Marker(symbol="circle", size=6)
    series.graphicalProperties = GraphicalProperties(
        ln=LineProperties(solidFill=PALETTE.accent, w=20000)
    )
    chart.series.append(series)
    return chart


# --------------------------------------------------------------------------- #
# supporting sheets
# --------------------------------------------------------------------------- #


def _bridge_sheet(ws: Worksheet, summary: PortfolioSummary) -> None:
    row = _banner(
        ws, "PnL Bridge",
        "Top to bottom from the model to the marks. Indented italics are 'of which' memo "
        "lines and are NOT part of any sum - which legs make up a row's duration total "
        "depends on that row's framework, so summing the leg columns across the book "
        "would double count.",
        width=5,
    )
    _table_header(ws, row, ["Component", "PnL", "% of practical", "Kind", "Why this line exists"])
    row += 1

    practical = summary.bridge.value("Practical PnL (bonds + linked hedges)")
    denom = abs(practical) if practical == practical and practical else float("nan")

    for line in summary.bridge.lines:
        label = ("    " + line.label) if line.is_memo else line.label
        cells = [
            ws.cell(row=row, column=1, value=label),
            ws.cell(row=row, column=2, value=_clean(line.amount)),
            ws.cell(row=row, column=3, value=_clean(line.amount / denom if denom == denom else None)),
            ws.cell(row=row, column=4, value=line.kind),
            ws.cell(row=row, column=5, value=line.note),
        ]
        cells[1].number_format = _EUR_C
        cells[2].number_format = _PCT
        cells[4].alignment = WRAP_LEFT
        if line.is_memo:
            for c in cells:
                c.font = body_font(9, italic=True, colour=PALETTE.memo)
        elif line.kind == "check":
            for c in cells:
                c.font = body_font(9, italic=True, colour=PALETTE.muted)
            cells[1].number_format = "0.000000"
            ok = abs(_clean(line.amount) or 0.0) < 1e-4
            cells[1].font = body_font(
                10, bold=True, colour=PALETTE.positive if ok else PALETTE.negative
            )
        elif line.kind in ("subtotal", "total"):
            for c in cells:
                c.font = body_font(11, bold=True, colour=PALETTE.ink)
                c.fill = fill(PALETTE.accent_soft)
        else:
            cells[0].font = body_font(10)
            cells[1].font = body_font(10, colour=_tone(line.amount))
        for c in cells:
            c.border = thin_border()
        ws.row_dimensions[row].height = 26
        row += 1

    _widths(ws, {1: 48, 2: 18, 3: 14, 4: 12, 5: 72})
    ws.sheet_view.showGridLines = False


def _snail_sheet(ws: Worksheet, trail: SnailTrail, summary: PortfolioSummary) -> None:
    row = _banner(
        ws, "Snail Trail",
        "A path, not a bar chart. The shape is the information: a tight coil round the "
        "origin is a model that tracks; a steady drift to one side is a systematic bias.",
        width=9,
    )

    row = _section(
        ws, row, "Today, walked leg by leg",
        "Start at zero, add one leg at a time. 'Still outstanding' is what the marks say "
        "is left to explain after that leg.",
        width=9,
    )
    _table_header(ws, row, ["#", "Leg", "Amount", "Cumulative explained", "Still outstanding", "Share of practical", "What this leg is"])
    row += 1
    for step in trail.steps:
        cells = [
            ws.cell(row=row, column=1, value=step.order),
            ws.cell(row=row, column=2, value=step.label),
            ws.cell(row=row, column=3, value=_clean(step.amount)),
            ws.cell(row=row, column=4, value=_clean(step.cumulative_explained)),
            ws.cell(row=row, column=5, value=_clean(step.outstanding)),
            ws.cell(row=row, column=6, value=_clean(step.share_of_practical)),
            ws.cell(row=row, column=7, value=step.note),
        ]
        for c, fmt in zip(cells, ["0", "@", _EUR, _EUR, _EUR, _PCT, "@"]):
            c.number_format = fmt
            c.font = body_font(10, bold=(step.label == "Residual"))
            c.border = thin_border()
        cells[6].alignment = WRAP_LEFT
        cells[6].font = body_font(9, italic=True, colour=PALETTE.muted)
        ws.row_dimensions[row].height = 26
        row += 1
    row += 2

    # ---- history ------------------------------------------------------------ #
    row = _section(
        ws, row, "The trail over time",
        "One point per run, cumulative. A perfect model traces the 45-degree line; the "
        "distance from it is the accumulated unexplained PnL.",
        width=9,
    )

    reading = trail.diagnostics
    verdict = str(reading.get("verdict", ""))
    ws.merge_cells(start_row=row, start_column=1, end_row=row, end_column=9)
    cell = ws.cell(row=row, column=1, value=f"Reading:  {verdict}")
    cell.font = body_font(11, bold=True, colour=PALETTE.ink)
    cell.fill = fill(PALETTE.accent_soft)
    cell.alignment = WRAP_LEFT
    ws.row_dimensions[row].height = 30
    row += 2

    for label, key, fmt in (
        ("Days of history", "days", "0"),
        ("Cumulative practical", "cum_practical", _EUR),
        ("Cumulative explained", "cum_explained", _EUR),
        ("Cumulative residual", "cum_residual", _EUR_C),
        ("Cumulative residual %", "cum_residual_pct", _PCT),
        ("Mean daily residual", "mean_daily_residual", _EUR_C),
        ("Std of daily residual", "std_daily_residual", _EUR_C),
        ("Standard error of the mean", "standard_error", _EUR_C),
        ("t-statistic (mean / standard error)", "t_statistic", "0.00"),
        ("Drift ratio (mean / std)", "drift_ratio", "0.000"),
        ("Share of days on the same side", "same_sign_share", _PCT),
        ("Days explained within 5%", "days_within_5pct", _PCT),
    ):
        if key not in reading:
            continue
        ws.cell(row=row, column=1, value=label).font = body_font(10)
        value_cell = ws.cell(row=row, column=2, value=_clean(reading.get(key)))
        value_cell.number_format = fmt
        value_cell.font = body_font(10, bold=True)
        row += 1
    row += 2

    table_top = row
    _table_header(
        ws, row,
        ["Date", "Practical", "Explained", "Residual", "Cum practical", "Cum explained", "Cum residual", "Cum |DV01|"],
    )
    row += 1
    for point in trail.points:
        values = [
            point.date, point.practical, point.explained, point.residual,
            point.cum_practical, point.cum_explained, point.cum_residual, point.cum_abs_dv01,
        ]
        formats = ["dd/mm/yyyy", _EUR, _EUR, _EUR_C, _EUR, _EUR, _EUR_C, _EUR]
        for i, (value, fmt) in enumerate(zip(values, formats), start=1):
            cell = ws.cell(row=row, column=i, value=_clean(value))
            cell.number_format = fmt
            cell.font = body_font(10)
            cell.border = thin_border()
        row += 1
    table_bottom = row - 1

    _widths(ws, {1: 30, 2: 16, 3: 16, 4: 16, 5: 17, 6: 17, 7: 16, 8: 16, 9: 60})

    if trail.has_history:
        _safe_chart(
            lambda: _history_snail_chart(
                ws, table_top + 1, table_bottom, 5, 6,
                "Cumulative explained against cumulative practical",
            ),
            ws, f"J{table_top}", "history snail",
        )
        _safe_chart(
            lambda: _residual_vs_risk_chart(ws, table_top + 1, table_bottom),
            ws, f"J{table_top + 22}", "residual against risk",
        )
    else:
        note = ws.cell(
            row=table_top,
            column=10,
            value="The trail needs at least two runs. Run another date and it appears here.",
        )
        note.font = body_font(10, italic=True, colour=PALETTE.muted)

    ws.sheet_view.showGridLines = False


def _residual_vs_risk_chart(ws: Worksheet, top: int, bottom: int):
    """Cumulative residual against cumulative risk taken.

    The risk/return reading of the same trail: is the unexplained PnL growing
    with the risk being run, or independently of it?  A line that rises with
    risk is a model that scales badly; one that rises regardless is a bias.
    """
    chart = ScatterChart()
    chart.title = "Cumulative residual against cumulative risk taken"
    chart.height = 10
    chart.width = 18
    chart.x_axis.title = "Cumulative |DV01| deployed"
    chart.y_axis.title = "Cumulative residual"
    chart.x_axis.numFmt = "#,##0"
    chart.y_axis.numFmt = "#,##0"
    chart.x_axis.delete = False
    chart.y_axis.delete = False

    series = Series(
        Reference(ws, min_col=7, min_row=top, max_row=bottom),
        Reference(ws, min_col=8, min_row=top, max_row=bottom),
        title="residual",
    )
    series.marker = Marker(symbol="circle", size=6)
    series.graphicalProperties = GraphicalProperties(
        ln=LineProperties(solidFill=PALETTE.warn, w=20000)
    )
    chart.series.append(series)
    return chart


def _asset_class_sheet(ws: Worksheet, summary: PortfolioSummary) -> None:
    row = _banner(
        ws, "Instrument Types Compared",
        "Bonds, futures and swaps each measured on their own terms - the view the "
        "bond-row-only legacy report could not produce, because a hedge with no linked "
        "bond simply did not appear anywhere.",
        width=12,
    )
    _table_header(
        ws, row,
        ["Instrument type", "Positions", "Linked", "Unlinked", "Exposure", "DV01",
         "Theoretical", "Practical", "Basis", "Difference", "Unlinked PnL", "Coverage"],
    )
    row += 1
    for ac in summary.asset_classes:
        values = [
            ac.name, ac.count, ac.linked_count, ac.unlinked_count, ac.exposure, ac.dv01,
            ac.theoretical, ac.practical, ac.basis, ac.residual, ac.unlinked_practical, ac.coverage,
        ]
        formats = ["@", "0", "0", "0", _EUR, _EUR, _EUR, _EUR, _EUR_C, _EUR_C, _EUR_C, _PCT]
        for i, (value, fmt) in enumerate(zip(values, formats), start=1):
            cell = ws.cell(row=row, column=i, value=_clean(value))
            cell.number_format = fmt
            cell.font = body_font(10, colour=_tone(value) if i in (7, 8, 9, 10) else PALETTE.text)
            cell.border = thin_border()
        row += 1
    row += 1

    ws.cell(row=row, column=1, value="Notes").font = body_font(11, bold=True, colour=PALETTE.ink)
    row += 1
    for ac in summary.asset_classes:
        ws.cell(row=row, column=1, value=ac.name).font = body_font(10, bold=True)
        ws.merge_cells(start_row=row, start_column=2, end_row=row, end_column=12)
        note = ws.cell(row=row, column=2, value=ac.note)
        note.font = body_font(9, italic=True, colour=PALETTE.muted)
        note.alignment = WRAP_LEFT
        ws.row_dimensions[row].height = 30
        row += 1

    _widths(ws, {1: 26, 2: 10, 3: 9, 4: 10, 5: 18, 6: 15, 7: 16, 8: 16, 9: 15, 10: 15, 11: 15, 12: 11})
    ws.sheet_view.showGridLines = False


def _framework_sheet(ws: Worksheet, summary: PortfolioSummary) -> None:
    from ..frameworks import CHAINS

    row = _banner(
        ws, "Spread Frameworks",
        "A framework decides how the yield move is SPLIT. It never changes the total: every "
        "chain below is a telescoping identity in the quoted rates, which is why the same "
        "day can be reported several ways without any of them disagreeing.",
        width=8,
    )

    row = _section(
        ws, row, "How this book is being measured",
        "Weighted by DV01, not by row count. Ten small government lines and one enormous "
        "credit line is a book measured against swaps; a headcount would say the opposite.",
        width=8,
    )
    _table_header(ws, row, ["Framework", "Positions", "DV01", "DV01 share", "Duration PnL", "Ties exactly", "Chain", "What it means"])
    row += 1
    for entry in summary.framework_mix:
        values = [
            entry["framework"], entry["positions"], entry["dv01"], entry["dv01_share"],
            entry["duration_pnl"], "yes" if entry["ties_exactly"] else "approximately",
            entry["chain"], entry["meaning"],
        ]
        formats = ["@", "0", _EUR, _PCT, _EUR, "@", "@", "@"]
        for i, (value, fmt) in enumerate(zip(values, formats), start=1):
            cell = ws.cell(row=row, column=i, value=_clean(value))
            cell.number_format = fmt
            cell.font = body_font(10)
            cell.border = thin_border()
            if i == 8:
                cell.alignment = WRAP_LEFT
                cell.font = body_font(9, italic=True, colour=PALETTE.muted)
        ws.row_dimensions[row].height = 30
        row += 1
    row += 2

    row = _section(
        ws, row, "The chains in full",
        "G and I tie EXACTLY because G-spread and I-spread are defined as differences from "
        "the same yield. ASW, Z and OAS are quoted on their own conventions, so their chains "
        "tie approximately - and the identity check on the Bonds sheet is where that shows.",
        width=8,
    )
    _table_header(ws, row, ["Code", "Base curve", "Ties exactly", "Chain", "Credit leg", "Meaning"])
    row += 1
    for code, chain in CHAINS.items():
        values = [
            code, chain.base_curve or "-", "yes" if chain.exact else "approximately",
            " + ".join(chain.legs) if chain.legs else "-",
            chain.spread_leg or "-", chain.description,
        ]
        for i, value in enumerate(values, start=1):
            cell = ws.cell(row=row, column=i, value=value)
            cell.font = body_font(10 if i < 6 else 9, italic=(i == 6))
            cell.border = thin_border()
            if i == 6:
                cell.alignment = WRAP_LEFT
                cell.font = body_font(9, italic=True, colour=PALETTE.muted)
        ws.row_dimensions[row].height = 28
        row += 1

    _widths(ws, {1: 12, 2: 13, 3: 15, 4: 52, 5: 20, 6: 80, 7: 40, 8: 70})
    ws.sheet_view.showGridLines = False


def _breakdown_sheet(ws: Worksheet, summary: PortfolioSummary) -> None:
    row = _banner(
        ws, "Breakdowns",
        "The same bridge, cut by portfolio and by currency.",
        width=14,
    )
    headers = [
        "Positions", "MV (close)", "DV01", "Hedge DV01", "Residual DV01",
        "Duration", "Carry", "Convexity", "FX", "Hedge PnL",
        "Theoretical", "Explained", "Practical", "Residual",
    ]
    formats = ["0", _EUR, _EUR, _EUR, _EUR, _EUR, _EUR, _EUR, _EUR, _EUR, _EUR, _EUR, _EUR, _EUR_C]
    keys = [
        "positions", "mv_current", "dv01", "hedge_dv01", "residual_dv01",
        "duration_pnl", "carry", "convexity", "fx", "hedge_pnl",
        "theoretical_pnl", "total_explained", "practical_pnl", "residual_pnl",
    ]

    for title, rows, key in (
        ("By portfolio", summary.by_portfolio, "portfolio"),
        ("By currency", summary.by_currency, "currency"),
    ):
        row = _section(ws, row, title, width=15)
        _table_header(ws, row, [key.title()] + headers)
        row += 1
        for record in rows:
            ws.cell(row=row, column=1, value=str(record.get(key, ""))).font = body_font(10, bold=True)
            ws.cell(row=row, column=1).border = thin_border()
            for i, (k, fmt) in enumerate(zip(keys, formats), start=2):
                cell = ws.cell(row=row, column=i, value=_clean(record.get(k)))
                cell.number_format = fmt
                cell.font = body_font(10)
                cell.border = thin_border()
            row += 1
        row += 2

    _widths(ws, {1: 18, **{i: 15 for i in range(2, 16)}})
    ws.sheet_view.showGridLines = False


def _curves_sheet(ws: Worksheet, result: AttributionResult) -> None:
    row = _banner(
        ws, "Curves",
        "Every curve the run was calibrated on. Linear in tenor with FLAT extrapolation at "
        "both ends - extending the slope of the front segment back to a two-day bond can "
        "produce a negative discount factor on a steep curve.",
        width=8,
    )
    _table_header(ws, row, ["Currency", "Curve", "Snapshot", "Nodes", "Min tenor", "Max tenor", "Min rate %", "Max rate %"])
    row += 1
    for entry in result.inputs_curves_described:
        values = [
            entry["currency"], entry["curve_type"], entry["snapshot"], entry["nodes"],
            entry["min_tenor"], entry["max_tenor"], entry["min_rate_pct"], entry["max_rate_pct"],
        ]
        formats = ["@", "@", "@", "0", "0.000", "0.000", "0.0000", "0.0000"]
        for i, (value, fmt) in enumerate(zip(values, formats), start=1):
            cell = ws.cell(row=row, column=i, value=_clean(value))
            cell.number_format = fmt
            cell.font = body_font(10)
            cell.border = thin_border()
        row += 1
    row += 2

    ws.cell(row=row, column=1, value="Curve nodes").font = body_font(12, bold=True, colour=PALETTE.ink)
    row += 2
    _table_header(ws, row, ["Currency", "Curve", "Tenor (years)", "Rate T-1 %", "Rate T0 %", "Move (bp)"])
    row += 1
    for curve in sorted(result.curves, key=lambda c: (c.key.currency, c.key.curve_type, c.key.snapshot)):
        if curve.key.snapshot != "prior":
            continue
        current = result.curves.get(curve.key.currency, curve.key.curve_type, "current")
        for i, tenor in enumerate(curve.tenors.tolist()):
            prior_rate = float(curve.rates[i])
            current_rate = float(current.rate(np.array([tenor]))[0]) if current else float("nan")
            values = [
                curve.key.currency, curve.key.curve_type, tenor,
                prior_rate, current_rate, (current_rate - prior_rate) * 100.0,
            ]
            formats = ["@", "@", "0.0000", "0.0000", "0.0000", "0.00"]
            for j, (value, fmt) in enumerate(zip(values, formats), start=1):
                cell = ws.cell(row=row, column=j, value=_clean(value))
                cell.number_format = fmt
                cell.font = body_font(9)
                cell.border = thin_border()
            row += 1

    _widths(ws, {1: 12, 2: 12, 3: 15, 4: 14, 5: 14, 6: 13, 7: 14, 8: 14})
    ws.sheet_view.showGridLines = False


def _quality_sheet(ws: Worksheet, result: AttributionResult, summary: PortfolioSummary) -> None:
    q = summary.quality
    row = _banner(
        ws, "Data Quality",
        "What this report could NOT see, stated rather than implied. A total that quietly "
        "covers most of the book is the failure mode worth guarding against.",
        width=6,
    )

    row = _section(
        ws, row, "Coverage",
        "How much of the book produced a complete answer.", width=6,
    )
    for label, key, fmt in (
        ("Bond positions", "bond_count", "0"),
        ("Rows with no failing check", "rows_ok", "0"),
        ("Rows fully attributed", "rows_attributed", "0"),
        ("Rows with a practical PnL", "rows_with_practical", "0"),
        ("Share fully attributed", "coverage_explained", _PCT),
        ("Share with a practical PnL", "coverage_practical", _PCT),
    ):
        ws.cell(row=row, column=1, value=label).font = body_font(10)
        cell = ws.cell(row=row, column=2, value=_clean(q.get(key)))
        cell.number_format = fmt
        cell.font = body_font(10, bold=True)
        row += 1
    row += 2

    row = _section(
        ws, row, "What the bond rows cannot see",
        "Hedges attach to bonds by ISIN and nothing else. A hedge with no LinkedISIN sits "
        "on no bond row, so its PnL is absent from every bond-level total; and if one ISIN "
        "appears on two rows, BOTH claim its full hedge PnL and DV01. Neither can be fixed "
        "by arithmetic - they are data questions - so they are measured instead.",
        width=6,
    )
    for label, key, fmt in (
        ("Unlinked futures", "unlinked_futures_count", "0"),
        ("Unlinked futures PnL", "unlinked_futures_pnl", _EUR_C),
        ("Unlinked swaps", "unlinked_swap_count", "0"),
        ("Unlinked swap PnL", "unlinked_swap_pnl", _EUR_C),
        ("Hedges pointing at an unknown ISIN", "orphan_futures_count", "0"),
        ("  their PnL", "orphan_futures_pnl", _EUR_C),
        ("Swaps pointing at an unknown ISIN", "orphan_swap_count", "0"),
        ("  their PnL", "orphan_swap_pnl", _EUR_C),
        ("Bond rows sharing an ISIN", "shared_isin_rows", "0"),
        ("Rows whose duration chain does not tie", "identity_break_rows", "0"),
    ):
        ws.cell(row=row, column=1, value=label).font = body_font(10)
        cell = ws.cell(row=row, column=2, value=_clean(q.get(key)))
        cell.number_format = fmt
        value = _clean(q.get(key)) or 0
        cell.font = body_font(
            10, bold=True,
            colour=PALETTE.warn if isinstance(value, (int, float)) and value else PALETTE.text,
        )
        row += 1
    row += 2

    row = _section(ws, row, "Tie-outs", "All three must be zero.", width=6)
    for line in summary.bridge.lines:
        if line.kind != "check":
            continue
        ws.cell(row=row, column=1, value=line.label).font = body_font(10)
        cell = ws.cell(row=row, column=2, value=_clean(line.amount))
        cell.number_format = "0.00000000"
        ok = abs(_clean(line.amount) or 0.0) < 1e-4
        cell.font = body_font(10, bold=True, colour=PALETTE.positive if ok else PALETTE.negative)
        ws.cell(row=row, column=3, value="OK" if ok else "BREAK").font = body_font(
            10, bold=True, colour=PALETTE.positive if ok else PALETTE.negative
        )
        row += 1
    row += 2

    row = _section(ws, row, "Statuses", "First failing check per row, cause before symptom.", width=6)
    _table_header(ws, row, ["Status", "Rows"])
    row += 1
    statuses = q.get("status_counts", {})
    for status, count in sorted(statuses.items(), key=lambda kv: (-kv[1], kv[0])):
        ws.cell(row=row, column=1, value=status).font = body_font(
            10, bold=(status != "OK"), colour=PALETTE.text if status == "OK" else PALETTE.negative
        )
        ws.cell(row=row, column=2, value=int(count)).font = body_font(10)
        for c in (1, 2):
            ws.cell(row=row, column=c).border = thin_border()
        row += 1

    _widths(ws, {1: 46, 2: 20, 3: 12, 4: 14, 5: 14, 6: 14})
    ws.sheet_view.showGridLines = False


def _definitions_sheet(ws: Worksheet, frames: Mapping[str, ColumnFrame], result: AttributionResult) -> None:
    row = _banner(
        ws, "Definitions",
        "Every column, its unit, what it means economically and the formula behind it. "
        "Generated from the same schema the columns are built from, so it cannot drift "
        "out of step with the numbers.",
        width=6,
    )

    row = _section(
        ws, row, "Conventions",
        "The choices that change a number, stated in words rather than as configuration flags.",
        width=6,
    )
    from .json_out import _conventions

    for key, sentence in _conventions(result.config).items():
        ws.cell(row=row, column=1, value=key.replace("_", " ").title()).font = body_font(10, bold=True)
        ws.merge_cells(start_row=row, start_column=2, end_row=row, end_column=6)
        cell = ws.cell(row=row, column=2, value=sentence)
        cell.font = body_font(9, colour=PALETTE.text)
        cell.alignment = WRAP_LEFT
        ws.row_dimensions[row].height = 30
        row += 1
    row += 2

    for sheet_name, frame in frames.items():
        row = _section(ws, row, f"{sheet_name} columns", width=6)
        _table_header(ws, row, ["Group", "Column", "Label", "Unit", "Meaning", "Formula"])
        row += 1
        for entry in frame.describe_schema():
            values = [entry["group"], entry["column"], entry["label"], entry["unit"], entry["meaning"], entry["formula"]]
            for i, value in enumerate(values, start=1):
                cell = ws.cell(row=row, column=i, value=value)
                cell.font = body_font(9, bold=(i == 2))
                cell.border = thin_border()
                if i >= 5:
                    cell.alignment = WRAP_LEFT
                    cell.font = body_font(9, italic=(i == 6), colour=PALETTE.text if i == 5 else PALETTE.muted)
            ws.row_dimensions[row].height = 28
            row += 1
        row += 2

    _widths(ws, {1: 14, 2: 30, 3: 24, 4: 11, 5: 76, 6: 52})
    ws.sheet_view.showGridLines = False


# --------------------------------------------------------------------------- #
# entry point
# --------------------------------------------------------------------------- #


def write_workbook(
    result: AttributionResult,
    summary: PortfolioSummary,
    trail: SnailTrail,
    path: Path | None = None,
) -> Path:
    """Build the workbook and return where it was written."""
    cfg = result.config
    if path is None:
        filename = cfg.output.excel
        if cfg.output.date_stamp_filenames:
            stem, _, suffix = filename.rpartition(".")
            filename = f"{stem}_{cfg.run.as_of.isoformat()}.{suffix}"
        path = cfg.paths.output_dir / filename
    path.parent.mkdir(parents=True, exist_ok=True)

    wb = Workbook()
    wb.remove(wb.active)

    _dashboard(wb.create_sheet("Dashboard"), result, summary, trail)
    _bridge_sheet(wb.create_sheet("Bridge"), summary)
    _snail_sheet(wb.create_sheet("Snail"), trail, summary)
    _asset_class_sheet(wb.create_sheet("Asset Classes"), summary)

    _write_frame_sheet(
        wb.create_sheet("Bonds"), result.bonds,
        "Bond Attribution",
        "One row per bond position. Every column is defined on the Definitions sheet, and "
        "hovering a header shows its meaning and formula.",
        freeze_cols=2,
        highlight={
            "residual_pnl": "diverging",
            "dv01_current": "bar",
            "attribution_status": "status",
        },
    )
    _write_frame_sheet(
        wb.create_sheet("Futures"), result.futures,
        "Bond Futures",
        "One row per contract, measured against the government curve at its CTD's tenor. "
        "The gap between practical and theoretical is the CTD switch and the delivery option.",
        freeze_cols=1,
        highlight={"basis_pnl": "diverging", "status": "status"},
    )
    _write_frame_sheet(
        wb.create_sheet("Swaps"), result.swaps,
        "Interest-Rate Swaps",
        "One row per deal, measured against the curve its floating leg actually projects "
        "off. Synthetic rows are hedge TARGETS, not positions: they carry risk here and "
        "PnL nowhere.",
        freeze_cols=1,
        highlight={"basis_pnl": "diverging", "status": "status"},
    )

    _framework_sheet(wb.create_sheet("Frameworks"), summary)
    _breakdown_sheet(wb.create_sheet("Breakdowns"), summary)
    _curves_sheet(wb.create_sheet("Curves"), result)
    _quality_sheet(wb.create_sheet("Data Quality"), result, summary)
    _definitions_sheet(
        wb.create_sheet("Definitions"),
        {"Bonds": result.bonds, "Futures": result.futures, "Swaps": result.swaps},
        result,
    )

    wb.active = 0
    wb.save(path)
    return path
