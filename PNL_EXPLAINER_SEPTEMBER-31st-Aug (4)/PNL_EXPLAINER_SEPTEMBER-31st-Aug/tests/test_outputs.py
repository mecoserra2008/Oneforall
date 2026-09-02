"""Config, loaders, snail trail and the three deliverables.

The output tests care about two things a reader cannot check by eye: that the
JSON is valid strict JSON (no bare `NaN` tokens, which several parsers refuse),
and that re-running a date REPLACES its history row rather than adding a second
one - a trail with two points for the same day is wrong in a way that looks
entirely plausible.
"""

from __future__ import annotations

import datetime as _dt
import json
import shutil
import tempfile
import unittest
from dataclasses import replace
from pathlib import Path

import numpy as np
from openpyxl import load_workbook

from pnlx.aggregate import build_portfolio
from pnlx.cli import main
from pnlx.config import ConfigError, load_config
from pnlx.engine import AttributionEngine
from pnlx.loaders import load_inputs
from pnlx.report import append_history, write_csv_outputs, write_metadata, write_workbook
from pnlx.report.csv_out import history_row
from pnlx.snail import SnailTrail, build_history_snail, build_intraday_snail, load_history


class TestConfig(unittest.TestCase):
    def test_the_shipped_config_loads(self):
        config = load_config()
        self.assertEqual(config.run.days, (config.run.as_of - config.run.prior).days)
        self.assertGreater(config.run.days, 0)

    def test_a_backwards_period_is_refused(self):
        with tempfile.TemporaryDirectory() as tmp:
            path = Path(tmp) / "c.yaml"
            path.write_text("run:\n  prior: 2025-09-30\n  as_of: 2025-09-29\n")
            with self.assertRaises(ConfigError):
                load_config(path)

    def test_an_unknown_key_is_refused(self):
        with tempfile.TemporaryDirectory() as tmp:
            path = Path(tmp) / "c.yaml"
            path.write_text("run:\n  prior: 2025-09-29\n  as_of: 2025-09-30\n  wibble: 1\n")
            with self.assertRaises(ConfigError):
                load_config(path)

    def test_an_unknown_framework_override_is_refused(self):
        with tempfile.TemporaryDirectory() as tmp:
            path = Path(tmp) / "c.yaml"
            path.write_text(
                "run:\n  prior: 2025-09-29\n  as_of: 2025-09-30\n"
                "model:\n  spread_framework_override: NONSENSE\n"
            )
            with self.assertRaises(ConfigError):
                load_config(path)


class TestLoaders(unittest.TestCase):
    def test_the_sample_book_loads(self):
        config = load_config()
        inputs = load_inputs(config)
        summary = inputs.summary()
        self.assertGreater(summary["bonds"], 0)
        self.assertGreater(summary["futures"], 0)
        self.assertGreater(summary["swaps_plain"], 0)
        self.assertGreater(summary["swaps_synthetic"], 0)

    def test_the_base_currency_fx_is_pinned_to_one(self):
        config = load_config()
        inputs = load_inputs(config)
        self.assertEqual(inputs.fx[config.run.base_ccy], (1.0, 1.0))

    def test_isins_are_normalised_for_the_hedge_join(self):
        config = load_config()
        inputs = load_inputs(config)
        for isin in inputs.bonds.isin.tolist():
            self.assertEqual(isin, isin.strip().upper())


class TestSnail(unittest.TestCase):
    @classmethod
    def setUpClass(cls):
        cls.config = load_config()
        cls.result = AttributionEngine(cls.config, load_inputs(cls.config)).run()
        cls.summary = build_portfolio(cls.result)

    def test_the_walk_closes_on_the_practical_total(self):
        steps = build_intraday_snail(self.result.bonds, self.summary)
        self.assertAlmostEqual(steps[-1].outstanding, 0.0, places=6)
        self.assertEqual(steps[-1].label, "Residual")

    def test_each_step_carries_the_running_total_forward(self):
        steps = build_intraday_snail(self.result.bonds, self.summary)
        running = 0.0
        for step in steps[:-1]:
            running += step.amount
            self.assertAlmostEqual(step.cumulative_explained, running, places=6)

    def test_a_single_day_is_reported_as_no_trail_yet(self):
        points, reading = build_history_snail(
            [{"as_of": "2025-09-30", "practical_pnl": "100", "total_explained": "95",
              "residual_pnl": "5", "bond_dv01": "1000"}]
        )
        self.assertEqual(len(points), 1)
        self.assertIn("first day", reading["verdict"])

    def test_a_one_sided_residual_is_read_as_drift(self):
        rows = [
            {"as_of": f"2025-10-{d:02d}", "practical_pnl": "1000",
             "total_explained": "900", "residual_pnl": "100", "bond_dv01": "5000"}
            for d in range(1, 16)
        ]
        _, reading = build_history_snail(rows)
        self.assertIn("drift", reading["verdict"])

    def test_a_scattered_residual_is_usually_read_as_coiled(self):
        """Tested over many series, not one.

        The reading is a statistical statement, so a single seed can land on an
        unlikely draw and say nothing about whether the rule is right - seed 7
        gives a t of -2.4, which is a genuinely unusual sample and SHOULD read
        as something other than coiled.  What matters is the proportion: an
        unbiased series should be called coiled roughly two thirds of the time,
        which is what |t| < 1 means for a normal, and should essentially never
        be called a systematic drift.
        """
        coiled = drifting = 0
        trials = 120
        for seed in range(trials):
            rng = np.random.default_rng(seed)
            rows = [
                {"as_of": f"2025-10-{d:02d}", "practical_pnl": "1000",
                 "total_explained": "1000", "residual_pnl": str(float(rng.normal(0, 200))),
                 "bond_dv01": "5000"}
                for d in range(1, 26)
            ]
            verdict = build_history_snail(rows)[1]["verdict"]
            coiled += "coiled" in verdict
            drifting += "systematic drift" in verdict

        self.assertGreater(coiled / trials, 0.55)
        self.assertLess(drifting / trials, 0.05)

    def test_the_reading_is_sample_size_aware(self):
        """The same daily bias must read more strongly the longer it persists.

        A small one-sided residual over five days is not evidence; over two
        hundred it is.  If the reading did not scale with the sample, it would
        mean the same thing after a week as after a year, which is the mistake
        a raw mean-over-standard-deviation ratio makes.
        """
        def t_for(days: int) -> float:
            rng = np.random.default_rng(3)
            rows = [
                {"as_of": (_dt.date(2025, 1, 1) + _dt.timedelta(days=d)).isoformat(),
                 "practical_pnl": "1000", "total_explained": "900",
                 "residual_pnl": str(100.0 + float(rng.normal(0, 200))),
                 "bond_dv01": "5000"}
                for d in range(days)
            ]
            return abs(build_history_snail(rows)[1]["t_statistic"])

        self.assertLess(t_for(5), t_for(200))
        self.assertIn("systematic drift", _verdict_for_persistent_bias())

    def test_the_cumulative_columns_accumulate(self):
        rows = [
            {"as_of": "2025-10-01", "practical_pnl": "100", "total_explained": "90",
             "residual_pnl": "10", "bond_dv01": "1000"},
            {"as_of": "2025-10-02", "practical_pnl": "200", "total_explained": "180",
             "residual_pnl": "20", "bond_dv01": "1000"},
        ]
        points, _ = build_history_snail(rows)
        self.assertAlmostEqual(points[-1].cum_practical, 300.0)
        self.assertAlmostEqual(points[-1].cum_residual, 30.0)
        self.assertAlmostEqual(points[-1].cum_abs_dv01, 2000.0)


class TestDeliverables(unittest.TestCase):
    def setUp(self):
        self.tmp = Path(tempfile.mkdtemp())
        base = load_config()
        self.config = replace(base, paths=replace(base.paths, output_dir=self.tmp))
        self.result = AttributionEngine(self.config, load_inputs(self.config)).run()
        self.summary = build_portfolio(self.result)
        self.trail = SnailTrail(
            steps=build_intraday_snail(self.result.bonds, self.summary),
            points=[],
            diagnostics={"verdict": "test"},
        )

    def tearDown(self):
        shutil.rmtree(self.tmp, ignore_errors=True)

    def test_csv_outputs_have_a_row_per_position(self):
        written = write_csv_outputs(self.result, self.summary)
        for key, expected in (
            ("bonds", len(self.result.bonds)),
            ("futures", len(self.result.futures)),
            ("swaps", len(self.result.swaps)),
        ):
            lines = written[key].read_text(encoding="utf-8").strip().splitlines()
            self.assertEqual(len(lines) - 1, expected, key)

    def test_csv_blanks_missing_values_rather_than_zeroing_them(self):
        written = write_csv_outputs(self.result, self.summary)
        text = written["swaps"].read_text(encoding="utf-8")
        self.assertNotIn("nan", text.lower())
        self.assertIn(",,", text)  # at least one genuinely blank field

    def test_the_metadata_is_strict_json(self):
        path = write_metadata(self.result, self.summary, self.trail)
        raw = path.read_text(encoding="utf-8")
        # A bare NaN token parses in Python but is not JSON, and strict parsers
        # elsewhere refuse it.  Both checks matter.
        self.assertNotIn("NaN", raw)
        self.assertNotIn("Infinity", raw)
        document = json.loads(raw, parse_constant=_reject)
        self.assertEqual(document["schema"], "pnlx.daily_pnl_metadata/1")
        self.assertIn("conventions", document)
        self.assertIn("bridge", document)

    def test_the_metadata_records_every_convention_that_changes_a_number(self):
        path = write_metadata(self.result, self.summary, self.trail)
        conventions = json.loads(path.read_text(encoding="utf-8"))["conventions"]
        for key in ("risk_date", "coupon_carry", "funding", "fx", "convexity",
                    "pull_to_par", "futures_dv01", "swap_dv01", "signs", "units"):
            self.assertIn(key, conventions)
            self.assertTrue(conventions[key].strip())

    def test_history_replaces_a_rerun_rather_than_appending(self):
        append_history(self.result, self.summary)
        path, rows = append_history(self.result, self.summary)
        self.assertEqual(len(rows), 1)
        stamps = [r["as_of"] for r in load_history(path)]
        self.assertEqual(len(stamps), len(set(stamps)))

    def test_history_keeps_earlier_days(self):
        earlier = replace(self.config, run=replace(self.config.run,
                                                   prior=_dt.date(2025, 9, 26),
                                                   as_of=_dt.date(2025, 9, 29)))
        earlier_result = AttributionEngine(earlier, load_inputs(earlier)).run()
        append_history(earlier_result, build_portfolio(earlier_result))
        path, rows = append_history(self.result, self.summary)
        self.assertEqual(len(rows), 2)
        self.assertEqual([r["as_of"] for r in rows], ["2025-09-29", "2025-09-30"])

    def test_the_workbook_has_every_sheet_and_no_broken_chart(self):
        path = write_workbook(self.result, self.summary, self.trail)
        wb = load_workbook(path)
        for name in ("Dashboard", "Bridge", "Snail", "Asset Classes", "Bonds",
                     "Futures", "Swaps", "Frameworks", "Breakdowns", "Curves",
                     "Data Quality", "Definitions"):
            self.assertIn(name, wb.sheetnames)
        for name in wb.sheetnames:
            for row in wb[name].iter_rows():
                for cell in row:
                    if isinstance(cell.value, str):
                        self.assertNotIn("chart unavailable", cell.value)

    def test_the_workbook_carries_no_nan_cells(self):
        path = write_workbook(self.result, self.summary, self.trail)
        wb = load_workbook(path)
        for name in wb.sheetnames:
            for row in wb[name].iter_rows():
                for cell in row:
                    if isinstance(cell.value, float):
                        self.assertFalse(np.isnan(cell.value), f"{name}!{cell.coordinate}")

    def test_the_definitions_sheet_documents_every_bond_column(self):
        path = write_workbook(self.result, self.summary, self.trail)
        wb = load_workbook(path)
        documented = {
            row[1].value
            for row in wb["Definitions"].iter_rows(min_col=1, max_col=2)
            if row[1].value
        }
        for column in self.result.bonds.names:
            self.assertIn(column, documented, column)


class TestCli(unittest.TestCase):
    def test_a_full_run_exits_clean_and_the_tie_outs_hold(self):
        with tempfile.TemporaryDirectory() as tmp:
            code = main(["--output-dir", tmp, "--check", "--quiet"])
            self.assertEqual(code, 0)
            self.assertTrue((Path(tmp) / "pnl_explainer.xlsx").is_file())
            self.assertTrue((Path(tmp) / "daily_pnl_explained.csv").is_file())
            self.assertTrue((Path(tmp) / "daily_pnl_metadata.json").is_file())

    def test_no_excel_skips_only_the_workbook(self):
        with tempfile.TemporaryDirectory() as tmp:
            self.assertEqual(main(["--output-dir", tmp, "--no-excel", "--quiet"]), 0)
            self.assertFalse((Path(tmp) / "pnl_explainer.xlsx").exists())
            self.assertTrue((Path(tmp) / "daily_pnl_explained.csv").is_file())

    def test_forcing_a_framework_still_respects_per_bond_overrides(self):
        """`--framework` is the GLOBAL override, and a per-bond one outranks it.

        The precedence is per-bond file, then global, then automatic.  A
        per-bond override is a deliberate, persistent desk decision recorded
        against a named ISIN - a callable that must be measured on OAS stays on
        OAS whatever the run-level flag says.  If it did not, a routine
        run-wide switch would silently undo a decision someone made for a
        reason, which is precisely what an override exists to prevent.
        """
        config = load_config()
        overridden = set(load_inputs(config).override_codes)
        self.assertTrue(overridden, "sample book should carry per-bond overrides")

        with tempfile.TemporaryDirectory() as tmp:
            self.assertEqual(
                main(["--output-dir", tmp, "--framework", "OIS", "--no-excel", "--quiet"]), 0
            )
            text = (Path(tmp) / "daily_pnl_explained.csv").read_text(encoding="utf-8")
            header = text.splitlines()[0].split(",")
            isin_at = header.index("isin")
            framework_at = header.index("spread_framework")

            for line in text.splitlines()[1:]:
                fields = line.split(",")
                isin, framework = fields[isin_at], fields[framework_at]
                if isin in overridden:
                    self.assertEqual(framework, load_inputs(config).override_codes[isin], isin)
                else:
                    self.assertEqual(framework, "OIS", isin)

    def test_a_bad_date_is_a_configuration_error_not_a_crash(self):
        with tempfile.TemporaryDirectory() as tmp:
            code = main(["--output-dir", tmp, "--as-of", "2020-01-01", "--quiet"])
            self.assertEqual(code, 2)


def _verdict_for_persistent_bias() -> str:
    """A small one-sided residual, held for a year: the reading must catch it."""
    rng = np.random.default_rng(11)
    rows = [
        {"as_of": (_dt.date(2025, 1, 1) + _dt.timedelta(days=d)).isoformat(),
         "practical_pnl": "1000", "total_explained": "940",
         "residual_pnl": str(60.0 + float(rng.normal(0, 150))),
         "bond_dv01": "5000"}
        for d in range(200)
    ]
    return build_history_snail(rows)[1]["verdict"]


def _reject(token):  # pragma: no cover - only fires on a malformed file
    raise ValueError(f"non-JSON constant in metadata: {token}")


if __name__ == "__main__":
    unittest.main()
