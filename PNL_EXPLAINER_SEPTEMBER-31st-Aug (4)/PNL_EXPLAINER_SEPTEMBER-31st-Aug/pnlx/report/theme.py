"""One visual language for the workbook.

Every colour, font, border and number format used by the report is declared
here, so the sheets read as one document rather than as twelve separately
formatted tables.  It also means a house style is one file to change.

The number formats matter more than the colours.  A PnL column shown to six
decimal places is unreadable and a basis-point column shown to none is wrong,
so the unit tag on each `Column` picks its format automatically and no sheet
gets to decide for itself.
"""

from __future__ import annotations

from dataclasses import dataclass
from typing import Final

from openpyxl.styles import Alignment, Border, Font, PatternFill, Side

from ..columns import Kind

__all__ = [
    "PALETTE",
    "FORMATS",
    "title_font",
    "header_font",
    "body_font",
    "fill",
    "thin_border",
    "bottom_rule",
    "number_format_for",
    "COLUMN_WIDTH",
]


# --------------------------------------------------------------------------- #
# palette
# --------------------------------------------------------------------------- #


@dataclass(frozen=True, slots=True)
class Palette:
    ink: str = "1B2A41"          # headline bars, section titles
    ink_soft: str = "2E4058"
    surface: str = "FFFFFF"
    panel: str = "F4F6F9"        # tile and table backgrounds
    panel_alt: str = "EAEEF4"
    rule: str = "D8DEE7"         # gridlines and borders
    text: str = "1F2933"
    muted: str = "6B7684"        # captions, notes
    positive: str = "1E7A4B"
    negative: str = "B3261E"
    warn: str = "B26A00"
    accent: str = "2A6FB5"
    accent_soft: str = "DCE8F5"
    memo: str = "8A94A6"


PALETTE: Final[Palette] = Palette()


# --------------------------------------------------------------------------- #
# number formats
# --------------------------------------------------------------------------- #

#: Negative amounts are shown in red and in parentheses.  A minus sign is easy
#: to miss at the end of a long column of figures, and a PnL sign read wrongly
#: is the most expensive kind of misread there is.
_EUR = '#,##0;[Red](#,##0)'
_EUR_CENTS = '#,##0.00;[Red](#,##0.00)'

FORMATS: Final[dict[str, str]] = {
    Kind.TEXT: "@",
    Kind.DATE: "dd/mm/yyyy",
    Kind.EUR: _EUR,
    Kind.BP: '0.00;[Red]-0.00',
    Kind.PCT: '0.0%;[Red]-0.0%',
    Kind.RATIO: '0.000;[Red]-0.000',
    Kind.RATE: '0.0000',
    Kind.PRICE: '0.0000',
    Kind.NUM: '#,##0.00',
    Kind.COUNT: '0',
    Kind.NOTIONAL: '#,##0;[Red](#,##0)',
    Kind.FLAG: "@",
}

#: Widths that fit the content without a manual pass.  Text columns get more
#: room because ISINs and status sentences are the ones people actually read.
COLUMN_WIDTH: Final[dict[str, int]] = {
    Kind.TEXT: 22,
    Kind.DATE: 12,
    Kind.EUR: 15,
    Kind.BP: 11,
    Kind.PCT: 11,
    Kind.RATIO: 11,
    Kind.RATE: 11,
    Kind.PRICE: 12,
    Kind.NUM: 13,
    Kind.COUNT: 9,
    Kind.NOTIONAL: 16,
    Kind.FLAG: 9,
}


def number_format_for(kind: str, *, cents: bool = False) -> str:
    if cents and kind == Kind.EUR:
        return _EUR_CENTS
    return FORMATS.get(kind, "General")


# --------------------------------------------------------------------------- #
# type
# --------------------------------------------------------------------------- #


def title_font(size: int = 18, colour: str | None = None) -> Font:
    return Font(name="Calibri", size=size, bold=True, color=colour or PALETTE.surface)


def header_font(size: int = 10, colour: str | None = None, bold: bool = True) -> Font:
    return Font(name="Calibri", size=size, bold=bold, color=colour or PALETTE.surface)


def body_font(
    size: int = 10,
    *,
    bold: bool = False,
    italic: bool = False,
    colour: str | None = None,
) -> Font:
    return Font(
        name="Calibri",
        size=size,
        bold=bold,
        italic=italic,
        color=colour or PALETTE.text,
    )


def fill(colour: str) -> PatternFill:
    return PatternFill(fill_type="solid", start_color=colour, end_color=colour)


def thin_border(colour: str | None = None) -> Border:
    side = Side(style="thin", color=colour or PALETTE.rule)
    return Border(left=side, right=side, top=side, bottom=side)


def bottom_rule(colour: str | None = None, style: str = "thin") -> Border:
    return Border(bottom=Side(style=style, color=colour or PALETTE.rule))


CENTRE: Final[Alignment] = Alignment(horizontal="center", vertical="center", wrap_text=True)
LEFT: Final[Alignment] = Alignment(horizontal="left", vertical="center")
RIGHT: Final[Alignment] = Alignment(horizontal="right", vertical="center")
WRAP_LEFT: Final[Alignment] = Alignment(horizontal="left", vertical="top", wrap_text=True)
